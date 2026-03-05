#!/usr/bin/env python3
import openpyxl
import json
import os

file_path = r'D:\your-mcp-proxy\AITY_VIP\docs\api\打板设计文档V1.3_何俊锋_20250306.xlsx'

try:
    if not os.path.exists(file_path):
        print(f"文件不存在: {file_path}")
        exit(1)

    print(f"正在读取文件: {file_path}\n")

    wb = openpyxl.load_workbook(file_path, data_only=False)

    print(f"\n成功读取Excel文件")
    print(f"工作表列表: {wb.sheetnames}\n")

    all_api_docs = {}

    for sheet_name in wb.sheetnames:
        print(f"\n{'='*80}")
        print(f"工作表: {sheet_name}")
        print(f"{'='*80}")

        ws = wb[sheet_name]
        print(f"维度: {ws.dimensions}")
        print(f"行数: {ws.max_row}, 列数: {ws.max_column}\n")

        # 读取内容
        print(f"\n内容:")
        for row_idx in range(1, min(ws.max_row + 1, 50)):
            row_data = []
            for col_idx in range(1, min(ws.max_column + 1, 10)):
                cell = ws.cell(row=row_idx + 1, column=col_idx)
                value = cell.value
                if value and not isinstance(cell, openpyxl.cell.cell.Cell):
                    continue
                if value:
                    row_data.append(str(value)[:100])
            if row_data:
                print(f"第{row_idx}行: {' | '.join(row_data)}")

        print(f"\n")

except Exception as e:
    print(f"\n错误: {e}")
    import traceback
    traceback.print_exc()
